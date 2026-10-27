from __future__ import annotations

import argparse
import hashlib
import json
import logging
import re
import sys
import zipfile
from pathlib import Path


MAX_BYTES = 50 * 1024 * 1024
MAX_PREVIEW_CHARS = 1600
MAX_SENSITIVE_SCAN_CHARS = 2_000_000
MAX_OOXML_EXPANDED_BYTES = 200 * 1024 * 1024
MAX_OOXML_MEMBERS = 10_000
ALLOWED = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md", ".csv", ".json"}

SENSITIVE_RULES = [
    ("身份证号", re.compile(r"(?<!\d)\d{6}(?:19|20)\d{2}[01]\d[0-3]\d\d{3}[0-9Xx](?!\d)")),
    ("手机号", re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")),
    ("邮箱", re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b")),
    ("银行卡号样式", re.compile(r"(?<!\d)(?:\d[ -]?){16,19}(?!\d)")),
    ("长交易识别号", re.compile(r"(?<!\d)\d{20,40}(?!\d)")),
    (
        "微信号",
        re.compile(
            r"(?im)(?:^微信号\s*$\n\s*[A-Za-z0-9_*.-]{3,64}\s*$|\bwxid_[a-z0-9_-]{6,}\b)"
        ),
    ),
    ("中文姓名字段", re.compile(r"(?m)^姓名\s*\n?\s*[\u4e00-\u9fff·]{2,20}\s*$")),
    ("财务凭证字段", re.compile(r"(?:转账电子凭证|交易金额|付款方|收款方|银行卡号|开户行)")),
]


def redact(text: str) -> tuple[str, list[str]]:
    # Detect every category against the original text.  Running detection on
    # already substituted markers can create false positives (for example the
    # marker for a bank-card pattern itself contains the word "银行卡号").
    original = text
    matched = [(label, pattern) for label, pattern in SENSITIVE_RULES if pattern.search(original)]
    findings: set[str] = {label for label, _ in matched}
    for label, pattern in matched:
        text = pattern.sub(f"【已隐藏{label}】", text)
    return text, sorted(findings)


def guard_ooxml(path: Path) -> None:
    if path.suffix.lower() not in {".docx", ".pptx", ".xlsx"}:
        return
    if not zipfile.is_zipfile(path):
        raise SystemExit("Office 文件不是有效的 OOXML 压缩包")
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > MAX_OOXML_MEMBERS:
            raise SystemExit("Office 文件内部成员过多，已拒绝解析")
        expanded = sum(max(0, item.file_size) for item in members)
        if expanded > MAX_OOXML_EXPANDED_BYTES:
            raise SystemExit("Office 文件解压后超过 200 MiB 安全上限")
        for item in members:
            if item.compress_size == 0 and item.file_size > 0:
                raise SystemExit("Office 文件包含异常压缩成员，已拒绝解析")
            if item.compress_size and item.file_size / item.compress_size > 200:
                raise SystemExit("Office 文件压缩比异常，已拒绝解析")


def read_text(path: Path) -> tuple[str, dict]:
    suffix = path.suffix.lower()
    meta: dict = {}
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        meta["pages"] = len(reader.pages)
        return "\n".join(page.extract_text() or "" for page in reader.pages), meta
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        meta["paragraphs"] = len(doc.paragraphs)
        return "\n".join(p.text for p in doc.paragraphs), meta
    if suffix == ".pptx":
        from pptx import Presentation

        deck = Presentation(str(path))
        meta["slides"] = len(deck.slides)
        chunks: list[str] = []
        for index, slide in enumerate(deck.slides, 1):
            chunks.append(f"[第 {index} 页]")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    chunks.append(shape.text)
        return "\n".join(chunks), meta
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(str(path), read_only=True, data_only=True)
        meta["sheets"] = book.sheetnames
        chunks = []
        for sheet in book.worksheets:
            chunks.append(f"[工作表 {sheet.title}]")
            for row in sheet.iter_rows(max_row=100, max_col=30, values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    chunks.append("\t".join(values))
        return "\n".join(chunks), meta
    return path.read_text(encoding="utf-8", errors="replace"), meta


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description="Safely inspect a document and emit redacted JSON.")
    parser.add_argument("path")
    args = parser.parse_args()
    path = Path(args.path).resolve()
    if not path.is_file():
        raise SystemExit("文件不存在")
    if path.suffix.lower() not in ALLOWED:
        raise SystemExit(f"不支持的文件类型：{path.suffix or '(无扩展名)'}")
    size = path.stat().st_size
    if size > MAX_BYTES:
        raise SystemExit(f"文件超过 {MAX_BYTES // 1024 // 1024} MiB 安全上限")

    guard_ooxml(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    text, meta = read_text(path)
    normalized = re.sub(r"[ \t]+", " ", text)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized).strip()
    _, findings = redact(normalized[:MAX_SENSITIVE_SCAN_CHARS])
    safe_preview, _ = redact(normalized[:MAX_PREVIEW_CHARS])
    if findings:
        safe_preview = "检测到敏感字段，预览已整体隐藏；仅供本机授权处理。"
    result = {
        "name": path.name,
        "extension": path.suffix.lower(),
        "bytes": size,
        "sha256_prefix": digest[:12],
        "characters_extracted": len(normalized),
        "sensitive_patterns_detected": findings,
        "preview_redacted": safe_preview,
        **meta,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    logging.getLogger("pypdf").setLevel(logging.CRITICAL)
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except Exception as exc:
        # Do not emit a traceback or an absolute local path to messaging logs.
        print(
            json.dumps(
                {
                    "status": "error",
                    "error": "document_parse_failed",
                    "exception": type(exc).__name__,
                },
                ensure_ascii=False,
            ),
            file=sys.stderr,
        )
        raise SystemExit(2)
